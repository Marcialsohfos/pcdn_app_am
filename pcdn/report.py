"""Rapport de guidage pour les équipes de traitement (XLSX de suivi + HTML lisible / imprimable).

Le rapport décrit ce qu'il y a à corriger, où (zone / fichier / ligne / identifiant), pourquoi et comment — il ne
modifie jamais les données exportées, qui restent brutes.
"""
from __future__ import annotations

import html
import io
from datetime import datetime

import pandas as pd

from .config import LAYERS
from .exporters import shp_field_map, shp_notes
from .qa import duplicate_table, status_series

APP_VERSION = "2.0 (export brut)"

ETAPES = {
    1: "Sources, fichiers et couverture",
    2: "Géométries",
    3: "Identifiants",
    4: "Doublons",
    5: "Valeurs hors dictionnaire",
    6: "Types et plages de valeurs",
    7: "Champs obligatoires manquants",
    8: "Cohérence métier et inter-couches",
    9: "Informations (aucune action requise)",
}
PRIO = {"ERROR": "P1 – À corriger", "WARNING": "P2 – À vérifier", "INFO": "P3 – Information"}

# code : (intitulé, action recommandée, étape)
RULES = {
    "STRUCT_UNREADABLE": ("Fichier illisible", "Redemander le fichier complet à l'équipe de collecte. Ses données NE SONT PAS dans l'export.", 1),
    "STRUCT_DOWNLOAD_FAILED": ("Téléchargement échoué", "Relancer la récupération ou récupérer le fichier manuellement. Ses données NE SONT PAS dans l'export.", 1),
    "STRUCT_NOT_CLASSIFIED": ("Couche non identifiée", "Identifier la couche réelle du fichier. Ses données sont déjà exportées (couche « Non_classe__… ») : les rattacher à la bonne couche.", 1),
    "STRUCT_ZONE_MISSING": ("Zone non identifiée", "Déterminer la zone du levé (nom de fichier ou dossier « ZONE n ») et la renseigner dans la table de travail.", 1),
    "STRUCT_EMPTY_FILE": ("Fichier vide", "Demander à l'équipe de collecte si le levé est réellement vide ou si l'export a échoué.", 1),
    "STRUCT_NO_CRS": ("Système de coordonnées absent", "Confirmer le système d'origine (EPSG:4326 supposé). Si les points sont mal placés, le .prj manque.", 1),
    "STRUCT_BAD_COORDS": ("Coordonnées KML illisibles", "Demander un nouvel export KML ; des sommets ont été ignorés à la lecture.", 2),
    "STRUCT_SIDECAR_MISSING": ("Fichier associé du Shapefile manquant", "Redemander le jeu complet (.shp .shx .dbf .prj). Sans .dbf/.shx la lecture peut être incomplète.", 1),
    "STRUCT_MISSING_COL": ("Colonne du schéma absente du fichier", "Vérifier la version du formulaire de collecte. En KML, un champ jamais rempli est normalement absent.", 1),
    "STRUCT_AMBIGUOUS": ("Colonne tronquée ambiguë", "Identifier la colonne d'origine (le Shapefile limite les noms à 10 caractères) ; demander l'export KML/GPKG.", 2),
    "STRUCT_EXTRA_COL": ("Colonnes hors schéma", "Conservées dans l'export. Vérifier s'il s'agit de champs utiles ajoutés au formulaire.", 9),
    "STRUCT_RENAMED": ("Colonne reconnue sous un autre nom", "Aucune action : seul le nom de la colonne a été harmonisé, les valeurs sont intactes.", 9),
    "STRUCT_ID_FROM_NAME": ("ID absent mais <name> KML présent", "Le <name> peut servir d'identifiant ; à reporter dans le champ ID lors du traitement.", 9),
    "STRUCT_ZONE_FROM_ID": ("Zone déduite des identifiants", "Vérifier que la zone déduite est correcte.", 9),
    "STRUCT_REPROJECTED": ("Fichier reprojeté en EPSG:4326", "Aucune action : reprojection nécessaire à la fusion.", 9),
    "STRUCT_MULTIFILES": ("Plusieurs fichiers pour une même couche/zone", "Tous conservés. Voir les doublons probables (DUP_SRC) pour décider lesquels garder.", 9),
    "GEOM_NULL": ("Géométrie absente", "Retrouver la position (GPS, fiche, photo) ou renvoyer l'entité à l'enquêteur.", 2),
    "GEOM_TYPE": ("Type de géométrie inattendu", "Vérifier la couche de collecte et recréer l'entité avec le bon type.", 2),
    "GEOM_INVALID": ("Géométrie invalide", "Réparer la géométrie (QGIS : Vérifier la validité / Réparer les géométries) sur une copie de travail.", 2),
    "GEOM_ZERO": ("Coordonnées (0,0) : GPS non fixé", "Repositionner l'entité à partir de la fiche ou de la photo, ou refaire le relevé.", 2),
    "GEOM_BBOX": ("Hors du Cameroun", "Probable inversion latitude/longitude ou erreur GPS : corriger la position.", 2),
    "LINE_SHORT": ("Tronçon très court", "Vérifier s'il s'agit d'un tronçon réel ou d'un clic parasite.", 2),
    "ID_MISSING": ("Identifiant manquant", "Attribuer un identifiant selon la convention ZONE<n>XX<n> (continuer la numérotation de la zone).", 3),
    "ID_FORMAT": ("Format d'identifiant non conforme", "Remplacer par la valeur proposée (colonne « Valeur proposée ») après vérification.", 3),
    "ID_PREFIX": ("Préfixe d'identifiant ≠ couche", "L'identifiant appartient à une autre couche : vérifier que l'entité est dans la bonne couche.", 3),
    "ID_ZONE": ("Identifiant d'une autre zone", "Vérifier la zone du fichier ou l'identifiant (erreur de saisie ou fichier mal rangé).", 3),
    "ID_DUP": ("Identifiant en double", "Si les deux fiches sont identiques : doublon (voir DUP_SRC). Sinon renuméroter l'une d'elles.", 3),
    "ID_GAP": ("Trous dans la numérotation", "Aucune action obligatoire : vérifier que les numéros manquants ne correspondent pas à des fiches perdues.", 9),
    "DUP_SRC": ("Doublon probable (même ID, même position)", "Comparer les fiches ; garder la plus complète et écarter l'autre dans la table de travail (rien n'a été supprimé dans l'export).", 4),
    "GEOM_DUP": ("Points très proches", "Vérifier s'il s'agit d'un doublon réel (même objet relevé deux fois) avant toute fusion.", 4),
    "DOM_INVALID": ("Valeur absente du dictionnaire", "Remplacer par la valeur du dictionnaire (proposition si proche) ou, si c'est une vraie nouvelle valeur, la faire valider.", 5),
    "DOM_CASE": ("Casse / accents / tiret différents du dictionnaire", "Remplacer par la valeur canonique proposée (correction purement typographique).", 5),
    "MULTI_EXCL": ("« Aucun » combiné à d'autres choix", "Retirer « Aucun » ou les autres choix.", 5),
    "TYPE_INVALID": ("Valeur de type incorrect", "Corriger la valeur (nombre ou Oui/Non attendu) à partir de la fiche terrain.", 6),
    "TYPE_DECIMAL": ("Décimal à la place d'un entier", "Arrondir après vérification.", 6),
    "RANGE": ("Valeur hors plage plausible", "Vérifier l'unité et la saisie (ex. mètres/kilomètres, pourcentage).", 6),
    "REQ_MISSING": ("Champ obligatoire / recommandé non renseigné", "Compléter à partir de la fiche terrain ou renvoyer à l'enquêteur.", 7),
    "XL_GARE": ("Gare sans fiche correspondante dans l'autre couche", "Les formulaires Gares et Infrastructures doivent couvrir les mêmes gares : compléter la fiche manquante.", 8),
    "XL_DIST": ("Distance déclarée incohérente avec la position", "Vérifier la position de l'ouvrage ou la distance déclarée.", 8),
}
DEFAULT_L = ("Incohérence entre champs liés", "Vérifier les deux champs avec la fiche terrain ; corriger le champ erroné.", 8)


def rule_info(code: str, message: str = ""):
    if code in RULES:
        return RULES[code]
    if code.startswith("L_"):
        return (message or DEFAULT_L[0], DEFAULT_L[1], 8)
    return (code, "À analyser.", 8)


# ------------------------------------------------------------------ tables
def _examples(g: pd.DataFrame) -> str:
    vals = [str(x) for x in (g["id"].where(g["id"].notna(), g["uid"]).where(lambda s: s.notna(), g["fichier"])).dropna().unique()[:5]]
    return ", ".join(vals)


def plan_table(issues: pd.DataFrame) -> pd.DataFrame:
    if issues.empty:
        return pd.DataFrame(columns=["Priorité", "Étape", "Couche", "Règle", "Intitulé", "Entités", "Anomalies", "Zones", "Exemples", "Action recommandée"])
    rows = []
    for (layer, rule, sev), g in issues.groupby(["layer", "rule", "severity"], observed=True):
        title, action, etape = rule_info(rule, g["message"].iloc[0])
        n_ent = int(g["_rid"].nunique()) if g["_rid"].notna().any() else len(g)
        zones = ", ".join(sorted(str(z) for z in g["zone"].dropna().unique())) or "—"
        rows.append((PRIO[str(sev)], f"{etape} · {ETAPES[etape]}", layer, rule, title, n_ent, len(g), zones, _examples(g), action))
    df = pd.DataFrame(rows, columns=["Priorité", "Étape", "Couche", "Règle", "Intitulé", "Entités", "Anomalies", "Zones", "Exemples", "Action recommandée"])
    return df.sort_values(["Priorité", "Étape", "Entités"], ascending=[True, True, False], kind="stable").reset_index(drop=True)


def anomalies_table(issues: pd.DataFrame) -> pd.DataFrame:
    if issues.empty:
        return pd.DataFrame(columns=["Couche", "Zone", "UID", "ID", "Fichier", "Ligne", "Champ", "Valeur actuelle", "Règle", "Sévérité",
                                     "Message", "Valeur proposée", "Action recommandée", "Statut", "Correction appliquée", "Traité par", "Date"])
    d = issues.copy()
    d["Action recommandée"] = [rule_info(r, m)[1] for r, m in zip(d["rule"], d["message"])]
    d["Sévérité"] = d["severity"].astype(str)
    d = d.rename(columns={"layer": "Couche", "zone": "Zone", "uid": "UID", "id": "ID", "fichier": "Fichier", "ligne": "Ligne",
                          "field": "Champ", "valeur": "Valeur actuelle", "rule": "Règle", "message": "Message", "suggestion": "Valeur proposée"})
    d["Ligne"] = d["Ligne"].map(lambda v: None if pd.isna(v) else int(v))
    for c in ("Statut", "Correction appliquée", "Traité par", "Date"):
        d[c] = ""
    d["Statut"] = "À faire"
    order = {"ERROR": 0, "WARNING": 1, "INFO": 2}
    d["_o"] = d["Sévérité"].map(order)
    d = d.sort_values(["_o", "Couche", "Zone", "Règle", "UID"], na_position="last", kind="stable")
    return d[["Couche", "Zone", "UID", "ID", "Fichier", "Ligne", "Champ", "Valeur actuelle", "Règle", "Sévérité", "Message",
              "Valeur proposée", "Action recommandée", "Statut", "Correction appliquée", "Traité par", "Date"]].reset_index(drop=True)


def zone_table(layers: dict, issues: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for name, g in layers.items():
        st = status_series(issues, name)
        s = g["_rid"].map(st).fillna("OK") if len(st) else pd.Series("OK", index=g.index)
        z = g["pcdn_zone"].fillna("(zone ?)")
        for zone, idx in z.groupby(z).groups.items():
            n = len(idx)
            e, w = int((s.loc[idx] == "ERROR").sum()), int((s.loc[idx] == "WARNING").sum())
            rows.append((zone, name, n, e, w, round(100 * (n - e - w) / n, 1) if n else 0.0))
    return pd.DataFrame(rows, columns=["Zone", "Couche", "Entités", "En erreur (P1)", "À vérifier (P2)", "Conformes (%)"]).sort_values(
        ["Zone", "Couche"]).reset_index(drop=True)


def coverage_table(layers: dict, expected_zones: list[str] | None = None) -> pd.DataFrame:
    zones = sorted({z for g in layers.values() for z in g["pcdn_zone"].dropna().unique()} | set(expected_zones or []))
    names = list(dict.fromkeys(list(LAYERS) + list(layers)))
    rows = {n: {z: (int((layers[n]["pcdn_zone"] == z).sum()) if n in layers else 0) for z in zones} for n in names}
    df = pd.DataFrame(rows).T
    df.insert(0, "Total", [len(layers[n]) if n in layers else 0 for n in names])
    unk = [int(layers[n]["pcdn_zone"].isna().sum()) if n in layers else 0 for n in names]
    if any(unk):
        df["(zone ?)"] = unk
    df.index.name = "Couche"
    return df.reset_index()


def reconciliation_table(layers: dict, inventory: pd.DataFrame, verify: pd.DataFrame | None) -> pd.DataFrame:
    read = inventory[inventory.statut.str.startswith("✔")].groupby("couche")["entites"].sum()
    base = pd.DataFrame({"Couche": list(layers), "Lues dans les sources": [int(read.get(n, 0)) for n in layers],
                         "Fusionnées": [len(g) for g in layers.values()]})
    base["Sources → fusion"] = ["✅ identique" if a == b else "❌ ÉCART" for a, b in zip(base["Lues dans les sources"], base["Fusionnées"])]
    if verify is not None and len(verify):
        v = verify.drop(columns=["Fusionné"], errors="ignore").rename(columns={"Contrôle": "Relecture des exports"})
        base = base.merge(v, on="Couche", how="left")
    return base


def synthesis_points(layers, inventory, issues, plan, coverage, recon, expected_zones) -> list[tuple[str, str]]:
    pts = []
    n_ent = sum(len(g) for g in layers.values())
    n_ok = int(inventory.statut.str.startswith("✔").sum())
    zones = sorted({z for g in layers.values() for z in g["pcdn_zone"].dropna().unique()})
    pts.append(("info", f"{n_ent} entités brutes issues de {n_ok} fichier(s), réparties en {len(layers)} couche(s) et {len(zones)} zone(s)"
                        f"{' (' + ', '.join(zones) + ')' if zones else ''}."))
    pts.append(("ok", "Données exportées BRUTES : parmi les fichiers intégrés, aucune entité n'est supprimée, aucune valeur corrigée, aucun "
                      "identifiant renuméroté. Seules 5 colonnes de traçabilité sont ajoutées (pcdn_uid, pcdn_zone, src_file, src_fmt, src_row)."))
    if recon is not None and "Relecture des exports" in recon:
        bad = recon[(recon["Relecture des exports"] == "❌ ÉCART") | (recon["Sources → fusion"] == "❌ ÉCART")]
        pts.append(("ok", "Réconciliation : toutes les couches sont identiques entre sources, fusion et exports relus.") if bad.empty else
                   ("err", f"Réconciliation : écart sur {', '.join(bad['Couche'])} — ne pas utiliser l'export avant vérification."))
    nonint = inventory[~inventory.statut.str.startswith("✔")]
    nonint = nonint[~nonint.statut.str.startswith("⏸")]
    if len(nonint):
        pts.append(("err", f"{len(nonint)} fichier(s) NON intégrés (illisibles ou téléchargement échoué) : {', '.join(nonint['fichier'].head(6))}"
                           f"{'…' if len(nonint) > 6 else ''}. Leurs données ne sont pas dans l'export."))
    excl = inventory[inventory.statut.str.startswith("⏸")]
    if len(excl):
        pts.append(("warn", f"{len(excl)} fichier(s) exclus manuellement à l'étape 1 : {', '.join(excl['fichier'].head(6))}."))
    nc = [n for n in layers if n not in LAYERS]
    if nc:
        pts.append(("warn", f"{len(nc)} fichier(s) de couche non identifiée exportés sous « Non_classe__… » : {', '.join(nc)}."))
    missing = [n for n in LAYERS if n not in layers]
    if missing:
        pts.append(("warn", f"Couches non reçues : {', '.join(missing)}."))
    zcols = [c for c in coverage.columns if c not in ("Couche", "Total", "(zone ?)")]
    holes = [(r["Couche"], z) for _, r in coverage.iterrows() if r["Couche"] in layers and r["Couche"] in LAYERS for z in zcols if r[z] == 0]
    if holes and zcols:
        pts.append(("warn", f"{len(holes)} combinaison(s) couche × zone sans aucune entité pour des couches pourtant reçues (voir « Couverture »)."))
    if expected_zones:
        absent = [z for z in expected_zones if z not in zones]
        if absent:
            pts.append(("err", f"Zones attendues sans aucune donnée : {', '.join(absent)}."))
    e1 = issues[issues["severity"].astype(str) == "ERROR"]
    n_feat = e1.dropna(subset=["_rid"]).drop_duplicates(["layer", "_rid"]).shape[0]
    n_struct = e1[e1["_rid"].isna()].drop_duplicates(["layer", "fichier", "rule"]).shape[0]
    pts.append(("err" if (n_feat or n_struct) else "ok",
                f"Priorité 1 : {n_feat} entité(s) avec au moins une erreur à corriger"
                f"{f', et {n_struct} problème(s) de fichier ou de structure (par fichier et type)' if n_struct else ''} "
                f"(sur {sum(len(g) for g in layers.values())} entités). Détail dans « Plan de traitement »."))
    dup = duplicate_table(layers)
    if len(dup):
        pts.append(("warn", f"{dup['groupe'].nunique()} groupe(s) de doublons probables ({len(dup)} entités) — conservés dans l'export, à arbitrer."))
    return pts


def build_tables(layers: dict, issues: pd.DataFrame, inventory: pd.DataFrame, completeness: pd.DataFrame,
                 log_lines: list[str], expected_zones: list[str] | None = None, verify: pd.DataFrame | None = None) -> dict:
    expected_zones = expected_zones or []
    cov = coverage_table(layers, expected_zones)
    plan = plan_table(issues)
    recon = reconciliation_table(layers, inventory, verify)
    t = {
        "points": synthesis_points(layers, inventory, issues, plan, cov, recon, expected_zones),
        "plan": plan,
        "anomalies": anomalies_table(issues),
        "zones": zone_table(layers, issues),
        "doublons": duplicate_table(layers),
        "couverture": cov,
        "sources": inventory,
        "reconciliation": recon,
        "completude": completeness,
        "champs_shp": shp_field_map(layers).query("tronque").drop(columns="tronque"),
        "limites_shp": shp_notes(layers),
        "journal": list(log_lines),
        "generated": datetime.now().strftime("%d/%m/%Y %H:%M"),
    }
    return t


# ------------------------------------------------------------------ XLSX
_SEV_FILL = {"ERROR": "F8D7DA", "WARNING": "FFF3CD", "INFO": "E2E3E5"}


def to_xlsx(t: dict) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    buf = io.BytesIO()
    readme = pd.DataFrame({"Rapport de traitement PCDN — " + t["generated"]: [
        "Ce classeur guide les équipes de traitement. Les données exportées (GPKG / SHP / KML) sont BRUTES : rien n'a été supprimé ni corrigé.",
        "Chaque anomalie renvoie à une entité par son UID (colonne pcdn_uid des exports), son fichier source et sa ligne dans ce fichier.",
        "",
        "Ordre de travail conseillé :",
        *[f"  {k}. {v}" for k, v in ETAPES.items() if k != 9],
        "",
        "Feuilles : Synthèse · Plan de traitement (quoi, combien, comment) · Par zone (répartition du travail) · Anomalies (liste détaillée à suivre) ·",
        "Doublons · Couverture · Sources · Réconciliation · Complétude · Champs SHP tronqués · Journal.",
        "Feuille « Anomalies » : renseigner Statut / Correction appliquée / Traité par / Date pour suivre l'avancement.",
        "P1 = à corriger · P2 = à vérifier · P3 = information.",
    ]})
    lvl = {"err": "ALERTE", "warn": "ATTENTION", "ok": "OK", "info": "INFO"}
    synth = pd.DataFrame([(lvl[lv], tx) for lv, tx in t["points"]], columns=["Niveau", "Constat"])
    sheets = [("Lisez-moi", readme), ("Synthèse", synth), ("Plan de traitement", t["plan"]), ("Par zone", t["zones"]),
              ("Anomalies", t["anomalies"]), ("Doublons", t["doublons"]), ("Couverture", t["couverture"]), ("Sources", t["sources"]),
              ("Réconciliation", t["reconciliation"]), ("Complétude", t["completude"]), ("Champs SHP tronqués", t["champs_shp"]),
              ("Limites SHP", pd.DataFrame({"Limite du format Shapefile": t["limites_shp"] or ["Aucune"]})),
              ("Journal", pd.DataFrame({"journal": t["journal"] or ["—"]}))]
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for name, df in sheets:
            df.to_excel(w, sheet_name=name, index=False)
        head = PatternFill("solid", fgColor="1F3A4D")
        for ws in w.book.worksheets:
            for c in ws[1]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), head
                c.alignment = Alignment(vertical="center", wrap_text=True)
            if ws.title != "Lisez-moi":
                ws.freeze_panes = "A2"
            if ws.max_row > 1 and ws.max_column > 1:
                ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col[:300] if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(70, max(10, width + 2))
            if ws.title in ("Lisez-moi", "Synthèse"):
                ws.column_dimensions["A"].width = 120 if ws.title == "Lisez-moi" else 14
                if ws.title == "Synthèse":
                    ws.column_dimensions["B"].width = 140
                for row in ws.iter_rows(min_row=2):
                    for c in row:
                        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws = w.book["Anomalies"]
        hdr = {c.value: c.column for c in ws[1]}
        if ws.max_row > 1:
            dv = DataValidation(type="list", formula1='"À faire,En cours,Corrigé,Rejeté (justifié),Sans objet"', allow_blank=True)
            ws.add_data_validation(dv)
            col = get_column_letter(hdr["Statut"])
            dv.add(f"{col}2:{col}{ws.max_row}")
            sc = hdr["Sévérité"]
            for r in range(2, ws.max_row + 1):
                v = ws.cell(r, sc).value
                if v in _SEV_FILL:
                    ws.cell(r, sc).fill = PatternFill("solid", fgColor=_SEV_FILL[v])
        for ws in (w.book["Plan de traitement"],):
            pc = {c.value: c.column for c in ws[1]}.get("Priorité")
            for r in range(2, ws.max_row + 1):
                v = str(ws.cell(r, pc).value or "")
                k = "ERROR" if v.startswith("P1") else "WARNING" if v.startswith("P2") else "INFO"
                ws.cell(r, pc).fill = PatternFill("solid", fgColor=_SEV_FILL[k])
        for name in ("Plan de traitement", "Anomalies"):
            for row in w.book[name].iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, str) and len(c.value) > 60:
                        c.alignment = Alignment(wrap_text=True, vertical="top")
    return buf.getvalue()


# ------------------------------------------------------------------ HTML
CSS = """
:root{--ink:#1d2a33;--mut:#5d6b75;--line:#d5dce1;--bg:#fff;--soft:#f4f7f9;--err:#b3261e;--errbg:#fbe9e7;--warn:#8a5a00;--warnbg:#fff4d6;--ok:#1b6b3a;--okbg:#e4f4ea;--info:#34506a;--infobg:#e8eef3}
@media (prefers-color-scheme:dark){:root{--ink:#e5eaee;--mut:#a3b0b9;--line:#33414b;--bg:#12181d;--soft:#1a2229;--err:#ff8a80;--errbg:#3b1f1d;--warn:#ffcc66;--warnbg:#3a2f12;--ok:#7fd6a0;--okbg:#15301f;--info:#9fc0dc;--infobg:#1c2a36}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 "Segoe UI",system-ui,-apple-system,Roboto,Arial,sans-serif}
main{max-width:1180px;margin:0 auto;padding:28px 22px 60px}
h1{font-size:26px;margin:0 0 4px;letter-spacing:-.01em}h2{font-size:19px;margin:34px 0 10px;padding-top:14px;border-top:2px solid var(--ink)}h3{font-size:15.5px;margin:22px 0 6px}
.sub{color:var(--mut);margin:0 0 18px}
.pt{padding:9px 13px;margin:6px 0;border-left:5px solid var(--line);background:var(--soft);border-radius:2px}
.pt.err{border-color:var(--err);background:var(--errbg)}.pt.warn{border-color:var(--warn);background:var(--warnbg)}.pt.ok{border-color:var(--ok);background:var(--okbg)}.pt.info{border-color:var(--info);background:var(--infobg)}
.wrap{overflow-x:auto;margin:8px 0 4px;border:1px solid var(--line);border-radius:4px}
table{border-collapse:collapse;width:100%;font-size:13px}th{background:var(--ink);color:var(--bg);text-align:left;padding:6px 8px;position:sticky;top:0;white-space:nowrap}
td{padding:5px 8px;border-top:1px solid var(--line);vertical-align:top}tr:nth-child(even) td{background:var(--soft)}
td.n{text-align:right;font-variant-numeric:tabular-nums}td.p1{background:var(--errbg)!important;color:var(--err);font-weight:600}td.p2{background:var(--warnbg)!important;color:var(--warn);font-weight:600}td.p3{color:var(--mut)}
td.z{color:var(--err);font-weight:600;background:var(--errbg)!important}
ol.steps{padding-left:20px}ol.steps li{margin:3px 0}code{background:var(--soft);padding:1px 5px;border-radius:3px;font-size:12.5px}
.note{color:var(--mut);font-size:13px}
@media print{main{max-width:none;padding:0}.wrap{overflow:visible;border:none}h2{break-after:avoid}tr{break-inside:avoid}th{position:static}}
"""


def _esc(v) -> str:
    return "" if v is None or (isinstance(v, float) and pd.isna(v)) else html.escape(str(v))


def _table(df: pd.DataFrame, max_rows: int | None = None, zero_red: list[str] | None = None) -> str:
    if df.empty:
        return '<p class="note">Aucune donnée.</p>'
    d = df.head(max_rows) if max_rows else df
    th = "".join(f"<th>{_esc(c)}</th>" for c in d.columns)
    body = []
    for _, r in d.iterrows():
        tds = []
        for c in d.columns:
            v = r[c]
            cls = ""
            if isinstance(v, (int, float)) and not isinstance(v, bool) and not pd.isna(v):
                cls = "n"
                if zero_red and c in zero_red and v == 0:
                    cls = "n z"
            if c == "Priorité":
                cls = "p1" if str(v).startswith("P1") else "p2" if str(v).startswith("P2") else "p3"
            if isinstance(v, float) and not pd.isna(v) and v == int(v):
                v = int(v)
            tds.append(f'<td class="{cls}">{_esc(v)}</td>')
        body.append("<tr>" + "".join(tds) + "</tr>")
    more = f'<p class="note">… {len(df) - max_rows} ligne(s) de plus : voir le classeur Excel.</p>' if max_rows and len(df) > max_rows else ""
    return f'<div class="wrap"><table><thead><tr>{th}</tr></thead><tbody>{"".join(body)}</tbody></table></div>{more}'


def to_html(t: dict, title: str = "Rapport de traitement — PCDN Corridor") -> str:
    plan, anom = t["plan"], t["anomalies"]
    P = []
    P.append(f"<h1>{_esc(title)}</h1><p class='sub'>Généré le {t['generated']} · application PCDN {APP_VERSION} · "
             "données exportées brutes — ce rapport est un guide, il ne modifie rien.</p>")
    P.append("<h2>À lire d'abord</h2>")
    P += [f'<div class="pt {lv}">{_esc(tx)}</div>' for lv, tx in t["points"]]

    P.append("<h2>Méthode de traitement recommandée</h2><ol class='steps'>")
    for k, v in ETAPES.items():
        if k == 9:
            continue
        sub = plan[plan["Étape"].str.startswith(f"{k} ·") & ~plan["Priorité"].str.startswith("P3")] if len(plan) else plan
        n = int(sub["Anomalies"].sum()) if len(sub) else 0
        P.append(f"<li><b>{_esc(v)}</b> — {n} anomalie(s) à corriger ou vérifier</li>")
    P.append("</ol><p class='note'>Corriger dans l'ordre : une géométrie ou un identifiant faux fausse les contrôles suivants (doublons, cohérence).</p>")

    P.append("<h2>Réconciliation sources → fusion → exports</h2>"
             "<p class='note'>Entités lues dans chaque fichier source = entités fusionnées = entités relues dans les fichiers exportés (effectif et identifiants pcdn_uid).</p>")
    P.append(_table(t["reconciliation"]))

    P.append("<h2>Couverture couche × zone</h2><p class='note'>Une case à 0 en rouge = levé manquant (à demander à l'équipe de collecte).</p>")
    zc = [c for c in t["couverture"].columns if c not in ("Couche", "Total", "(zone ?)")]
    P.append(_table(t["couverture"], zero_red=zc))

    P.append("<h2>Répartition du travail par zone</h2>")
    P.append(_table(t["zones"]))

    P.append("<h2>Plan de traitement par couche</h2><p class='note'>Une ligne = une règle de contrôle. « Entités » = nombre de fiches concernées. "
             "Le détail fiche par fiche (UID, fichier, ligne, valeur actuelle, valeur proposée) est dans la feuille « Anomalies » du classeur Excel.</p>")
    main_plan = plan[~plan["Priorité"].str.startswith("P3")] if len(plan) else plan
    if main_plan.empty:
        P.append('<div class="pt ok">Aucune anomalie P1/P2.</div>')
    for layer, g in main_plan.groupby("Couche", sort=True):
        P.append(f"<h3>{_esc(layer)}</h3>")
        P.append(_table(g.drop(columns=["Couche"]), max_rows=40))

    if len(t["doublons"]):
        P.append("<h2>Doublons probables (conservés dans l'export)</h2><p class='note'>Même identifiant et même position (à 0,1 m près). Comparer les fiches d'un groupe et arbitrer.</p>")
        P.append(_table(t["doublons"].drop(columns=["layer"]).assign(entre_fichiers=lambda d: d["entre_fichiers"].map({True: "oui", False: "non"})), max_rows=60))

    P.append("<h2>Fichiers sources</h2>")
    P.append(_table(t["sources"].drop(columns=["chemin"]), max_rows=200))

    P.append("<h2>Limites des formats d'export</h2><ul>")
    P.append("<li><b>GeoPackage</b> : format de référence, complet (aucune limite de nom ni de longueur).</li>")
    P.append("<li><b>Shapefile</b> : noms de champs limités à 10 caractères (correspondance : <code>correspondance_champs_shp.csv</code>), texte limité à 254 caractères.</li>")
    P.append("<li><b>KML</b> : tous les attributs en ExtendedData ; les champs vides sont omis.</li></ul>")
    if t["limites_shp"]:
        P.append("<h3>Points d'attention Shapefile pour ce jeu de données</h3><ul>" + "".join(f"<li>{_esc(x)}</li>" for x in t["limites_shp"][:30]) + "</ul>")

    used = sorted(set(plan["Règle"])) if len(plan) else []
    if used:
        P.append("<h2>Glossaire des règles rencontrées</h2>")
        gl = pd.DataFrame([(r, rule_info(r)[0], rule_info(r)[1]) for r in used], columns=["Règle", "Signification", "Que faire"])
        P.append(_table(gl))
    return f'<!doctype html><html lang="fr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{_esc(title)}</title><style>{CSS}</style></head><body><main>{"".join(P)}</main></body></html>'
